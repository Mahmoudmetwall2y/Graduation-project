from pathlib import Path
import re
import sys

import fitz
from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def inspect_pdf(path: Path) -> None:
    pdf = fitz.open(path)
    print(f"PDF\t{path.name}\tpages={pdf.page_count}")
    toc = pdf.get_toc(simple=True)
    if toc:
        print("PDF_BOOKMARKS")
        for level, title, page in toc:
            print(f"{level}\t{page}\t{title}")
    print("PDF_TOC_CANDIDATES")
    for number in range(min(pdf.page_count, 30)):
        text = pdf[number].get_text("text")
        if re.search(r"table\s+of\s+contents|contents", text, re.I):
            clean = "\n".join(line.strip() for line in text.splitlines() if line.strip())
            print(f"--- page {number + 1} ---")
            print(clean[:10000])


def inspect_docx(path: Path) -> None:
    doc = Document(path)
    headings = []
    nonempty = 0
    words = 0
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        nonempty += 1
        words += len(text.split())
        style = paragraph.style.name if paragraph.style else ""
        if style.lower().startswith("heading") or re.match(
            r"^(chapter\s+\d+|\d+(?:\.\d+){0,4}\s+|references$|appendix)",
            text,
            re.I,
        ):
            headings.append((index, style, text))
    print(
        f"DOCX\t{path.name}\tparagraphs={len(doc.paragraphs)}\t"
        f"nonempty={nonempty}\twords={words}\ttables={len(doc.tables)}\t"
        f"sections={len(doc.sections)}"
    )
    for index, style, text in headings:
        print(f"H\t{index}\t{style}\t{text}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("usage: inspect_book_assets.py <relative-path> [...]")
    for raw in sys.argv[1:]:
        path = (ROOT / raw).resolve()
        if path.suffix.lower() == ".pdf":
            inspect_pdf(path)
        elif path.suffix.lower() == ".docx":
            inspect_docx(path)
        else:
            raise SystemExit(f"unsupported file: {path}")


if __name__ == "__main__":
    main()
