from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Graduation Book Template - New.docx"
SOURCES = [
    ROOT / "docs" / "CHAPTER_1_INTRODUCTION.md",
    ROOT / "docs" / "CHAPTER_2_BACKGROUND_RELATED_WORK.md",
    ROOT / "docs" / "CHAPTER_3_SYSTEM_DESIGN_IMPLEMENTATION.md",
    ROOT / "docs" / "CHAPTER_4_RESULTS_DISCUSSION.md",
    ROOT / "docs" / "CHAPTER_5_CONCLUSION_FUTURE_WORK.md",
    ROOT / "docs" / "CHAPTER_6_REFERENCES.md",
    ROOT / "docs" / "APPENDICES.md",
]
OUTPUT = ROOT / "docs" / "AscultiCor_Graduation_Book_Template_Filled.docx"

NAVY = "17365D"
TEAL = "0F6B78"
LIGHT_TEAL = "DDEBF0"
LIGHT_BLUE = "EAF1F8"
LIGHT_GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = "808080", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.extend((begin, instr, separate, text, end))


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def force_style_font(style, name: str, size: float) -> None:
    """Set an explicit Word font and remove template theme-font overrides."""
    r_pr = style._element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in fonts.attrib:
            del fonts.attrib[key]
    style.font.name = name
    style.font.size = Pt(size)


def style_document(document: Document) -> None:
    styles = document.styles
    if "Table Grid" not in styles:
        style = styles.add_style("Table Grid", WD_STYLE_TYPE.TABLE)
        style.base_style = styles["Normal Table"]

    if "Numbered Body" not in styles:
        style = styles.add_style("Numbered Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Numbered Body"]
    force_style_font(style, "Times New Roman", 12)
    style.paragraph_format.left_indent = Cm(0.7)
    style.paragraph_format.first_line_indent = Cm(-0.7)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_after = Pt(0)

    normal = styles["Normal"]
    force_style_font(normal, "Times New Roman", 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, NAVY, 18, 12),
        "Heading 2": (16, NAVY, 12, 5),
        "Heading 3": (14, TEAL, 9, 3),
        "Heading 4": (12, TEAL, 7, 3),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        force_style_font(style, "Times New Roman", size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "Front Matter Title" not in styles:
        style = styles.add_style("Front Matter Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Front Matter Title"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(18)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_after = Pt(18)
    style.paragraph_format.page_break_before = True

    if "Figure Placeholder" not in styles:
        style = styles.add_style("Figure Placeholder", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Figure Placeholder"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    style.font.italic = True
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_together = True

    if "Figure Caption" not in styles:
        style = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Figure Caption"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    style.font.italic = True
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(10)
    style.paragraph_format.keep_together = True

    if "Table Caption" not in styles:
        style = styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Table Caption"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

    if "Reference Entry" not in styles:
        style = styles.add_style("Reference Entry", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Reference Entry"]
    force_style_font(style, "Times New Roman", 12)
    style.paragraph_format.left_indent = Cm(0.7)
    style.paragraph_format.first_line_indent = Cm(-0.7)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_after = Pt(0)


def configure_multilevel_numbering(document: Document) -> int:
    try:
        numbering_part = document.part.numbering_part
    except NotImplementedError:
        element = parse_xml(
            '<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        )
        numbering_part = NumberingPart(
            PackURI("/word/numbering.xml"), CT.WML_NUMBERING, element, document.part.package
        )
        document.part.relate_to(numbering_part, RT.NUMBERING)
    numbering = numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId"), "").isdigit()
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId"), "").isdigit()
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)

    formats = ("Chapter %1:", "%1.%2", "%1.%2.%3", "%1.%2.%3.%4")
    styles = ("Heading1", "Heading2", "Heading3", "Heading4")
    for level, (text, style_id) in enumerate(zip(formats, styles)):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), style_id)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        lvl.extend((start, num_fmt, p_style, lvl_text, suff))
        abstract.append(lvl)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(abstract)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))


def add_page_number_footer(document: Document) -> None:
    for section in document.sections:
        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("— ")
        add_field(paragraph, "PAGE", "1")
        paragraph.add_run(" —")


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    logo = ROOT / "Logo.jpeg"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(1.45))

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("[UNIVERSITY NAME]").bold = True
    p.runs[0].font.size = Pt(15)
    p = document.add_paragraph("[FACULTY / DEPARTMENT]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("ASCULTICOR")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("An AI-Powered IoT Platform for ECG and PCG Acquisition,\nAnalysis, and Remote Cardiac Monitoring")
    run.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.add_run("Graduation Project Book\n").bold = True
    p.add_run("Submitted in partial fulfillment of the requirements for [DEGREE NAME]")

    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    values = (("Prepared by", "[STUDENT NAMES AND IDS]"), ("Supervised by", "[SUPERVISOR NAME AND TITLE]"))
    for row, values_row in zip(table.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], LIGHT_BLUE)

    p = document.add_paragraph("Academic Year 2025–2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.runs[0].bold = True


def add_front_page(document: Document, title: str, paragraphs: list[str]) -> None:
    document.add_paragraph(title, style="Front Matter Title")
    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_front_matter(document: Document) -> None:
    document.add_paragraph("Approval Page", style="Front Matter Title")
    p = document.add_paragraph(
        "This graduation project book, entitled “AscultiCor: An AI-Powered IoT Platform for ECG and PCG "
        "Acquisition, Analysis, and Remote Cardiac Monitoring,” has been submitted to [FACULTY / DEPARTMENT] "
        "in partial fulfillment of the requirements for [DEGREE NAME]."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    table = document.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for row, values in zip(
        table.rows,
        (
            ("Role", "Name", "Signature / Date"),
            ("Project Supervisor", "[NAME AND TITLE]", "________________"),
            ("Examiner", "[NAME AND TITLE]", "________________"),
            ("Department Approval", "[NAME AND TITLE]", "________________"),
        ),
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
            set_cell_margins(cell)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True

    add_front_page(
        document,
        "Declaration",
        [
            "We declare that this graduation project book is our original work, except where sources are "
            "acknowledged by citation. The implementation, experiments, figures, and conclusions are presented "
            "within the stated educational and research scope. This work has not been submitted in the same form "
            "for another degree or academic award.",
            "Student names and signatures: [INSERT APPROVED NAMES, IDS, AND SIGNATURES]",
            "Date: [INSERT DATE]",
        ],
    )
    add_front_page(
        document,
        "Dedication",
        ["[Insert the team's approved dedication, or remove this page if a dedication is not required.]"],
    )
    add_front_page(
        document,
        "Acknowledgements",
        [
            "[Insert the approved acknowledgements for the project supervisor, faculty members, technical "
            "contributors, participating organizations, families, and project-team members. Confirm names and "
            "titles before final submission.]"
        ],
    )
    add_front_page(
        document,
        "Abstract",
        [
            "AscultiCor is an educational cardiac-monitoring prototype that integrates simultaneous "
            "electrocardiogram (ECG) and phonocardiogram (PCG) acquisition with Internet-of-Things communication, "
            "artificial-intelligence analysis, a multi-tenant web application, workflow automation, and cloud "
            "deployment controls. The sensing device combines an ESP32-WROOM-32 with an AD8232 single-lead ECG "
            "front end and a MAX9814 acoustic path. Its firmware implements dual-rate sampling, preflight checks, "
            "buffered binary MQTT transport, provisioning, session control, status reporting, and recovery. A "
            "custom printed circuit board incorporates USB Type-C input, lithium-ion charging and protection, load "
            "sharing, 3.3 V regulation, decoupling, and analog–digital layout considerations. The inference layer "
            "deploys an XGBoost PCG classifier using traditional features and YAMNet embeddings, together with an "
            "AuscultiCor v26 single-lead ECG model and an active PyTorch multi-head convolutional neural network "
            "for murmur characterization. The CNN accepts four mel-spectrogram channels representing aortic, "
            "mitral, pulmonary, and tricuspid positions and returns timing, shape, grading, pitch, quality, and "
            "location outputs. The Next.js and Supabase "
            "application manages users, organizations, patients, devices, sessions, recordings, predictions, "
            "reports, alerts, and audit information, while n8n workflows coordinate asynchronous reporting, "
            "monitoring, notification, and escalation. On the supplied held-out PCG test matrix of 1,651 examples, "
            "the XGBoost model achieved 83.65% accuracy, 82.33% Murmur recall, 64.60% Murmur precision, a Murmur "
            "area under the receiver-operating-characteristic curve of 0.915, and average precision of 0.831. Four "
            "delivered ECG case studies demonstrate structured model outputs but do not establish aggregate ECG "
            "accuracy. Six delivered CNN confusion matrices each contain 895 examples and yield matrix accuracies "
            "from 95.42% to 97.09%; raw predictions, split provenance, and runtime CNN case studies remain to be "
            "added. Static checks confirmed frontend type and lint quality, Python syntax validity, and all 17 "
            "model-registry tests; hardware "
            "measurement, configured service integration, workflow execution histories, deployment trials, and "
            "prospective clinical validation remain future work. The project demonstrates an inspectable "
            "end-to-end engineering architecture and is not a certified diagnostic device.",
            "Keywords—ECG, PCG, cardiac monitoring, Internet of Things, XGBoost, deep learning, MQTT, Supabase, n8n.",
        ],
    )

    document.add_paragraph("Table of Contents", style="Front Matter Title")
    p = document.add_paragraph()
    add_field(p, 'TOC \\o "1-4" \\h \\z \\u', "Right-click and select Update Field")

    document.add_paragraph("List of Figures", style="Front Matter Title")
    p = document.add_paragraph()
    add_field(p, 'TOC \\h \\z \\t "Figure Caption,1"', "Right-click and select Update Field")

    document.add_paragraph("List of Tables", style="Front Matter Title")
    p = document.add_paragraph()
    add_field(p, 'TOC \\h \\z \\t "Table Caption,1"', "Right-click and select Update Field")

    document.add_paragraph("List of Abbreviations", style="Front Matter Title")
    abbreviations = [
        ("ADC", "Analog-to-Digital Converter"), ("AGC", "Automatic Gain Control"),
        ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"),
        ("AUC", "Area Under the Curve"),
        ("BPM", "Beats Per Minute"), ("CNN", "Convolutional Neural Network"),
        ("ECG", "Electrocardiogram"), ("IoT", "Internet of Things"),
        ("LLM", "Large Language Model"), ("MQTT", "Message Queuing Telemetry Transport"),
        ("PCG", "Phonocardiogram"), ("PCB", "Printed Circuit Board"),
        ("PR", "Precision–Recall"), ("QRS", "QRS Complex"),
        ("RLS", "Row-Level Security"), ("ROC", "Receiver Operating Characteristic"),
        ("SNR", "Signal-to-Noise Ratio"), ("SVEB", "Supraventricular Ectopic Beat"),
        ("TLS", "Transport Layer Security"), ("VEB", "Ventricular Ectopic Beat"),
        ("VPS", "Virtual Private Server"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    table.rows[0].cells[0].text = "Abbreviation"
    table.rows[0].cells[1].text = "Meaning"
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for short, meaning in abbreviations:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = short, meaning
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            set_cell_margins(cell)


def strip_heading_number(text: str) -> str:
    text = re.sub(r"^Chapter\s+\d+\s*:\s*", "", text, flags=re.I)
    text = re.sub(r"^\d+(?:\.\d+){0,3}\s+", "", text)
    return text.strip()


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    data = [rows[0]] + rows[2:]
    table = document.add_table(rows=1, cols=len(data[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for index, value in enumerate(data[0]):
        table.rows[0].cells[index].text = value
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.size = Pt(9.5)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        set_cell_margins(cell)
    for row_index, row_data in enumerate(data[1:]):
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = value
            set_cell_margins(cells[index])
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if row_index % 2:
                set_cell_shading(cells[index], LIGHT_GRAY)
    document.add_paragraph()


def add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    pieces = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(piece)


def add_real_figure(document: Document, caption: str, image_path: str) -> None:
    resolved = (ROOT / image_path).resolve()
    if not resolved.exists():
        paragraph = document.add_paragraph(style="Figure Placeholder")
        paragraph.add_run(f"Missing project figure: {caption}\nExpected file: {image_path}")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(resolved), width=Inches(6.35))
    document.add_paragraph(caption, style="Figure Caption")


def add_figure_grid(document: Document, caption: str, items: list[tuple[str, str]]) -> None:
    """Insert a compact row of real result panels without altering the images."""
    table = document.add_table(rows=2, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, (image_path, label) in enumerate(items):
        cell = table.rows[0].cells[index]
        resolved = (ROOT / image_path).resolve()
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        if resolved.exists():
            paragraph.add_run().add_picture(str(resolved), width=Inches(1.95))
        else:
            paragraph.add_run(f"Missing: {image_path}")
        set_cell_margins(cell, top=30, start=30, bottom=20, end=30)

        label_cell = table.rows[1].cells[index]
        label_cell.text = label
        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_paragraph.paragraph_format.space_after = Pt(0)
        for run in label_paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            run.font.italic = True
        set_cell_margins(label_cell, top=10, start=30, bottom=20, end=30)
    document.add_paragraph(caption, style="Figure Caption")


def add_chapter_from_markdown(document: Document, num_id: int, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    in_references = False
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1)) - 1
            raw_title = heading.group(2)
            if re.match(r"^chapter\s+\d+\s+references$", raw_title, flags=re.I):
                in_references = True
                chapter_number = re.search(r"\d+", raw_title).group()
                paragraph = document.add_paragraph(
                    f"Provisional References Cited in Chapter {chapter_number}", style="Heading 2"
                )
                index += 1
                continue
            paragraph = document.add_paragraph(strip_heading_number(raw_title), style=f"Heading {level + 1}")
            if not re.match(r"^Appendix\s+[A-Z]:", raw_title, flags=re.I):
                apply_numbering(paragraph, num_id, level)
            index += 1
            continue

        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_markdown_table(document, rows)
            continue

        image = re.match(r"^!\[(.+)\]\((.+)\)$", line)
        if image:
            add_real_figure(document, image.group(1).strip(), image.group(2).strip())
            index += 1
            continue

        if line.startswith("[[FIGURE_GRID|") and line.endswith("]]" ):
            fields = line[2:-2].split("|")
            caption = fields[1].strip()
            items = []
            for field in fields[2:]:
                image_path, label = field.split("::", 1)
                items.append((image_path.strip(), label.strip()))
            add_figure_grid(document, caption, items)
            index += 1
            continue

        if line.startswith("> **Figure"):
            paragraph = document.add_paragraph(style="Figure Placeholder")
            paragraph.add_run("\n" + re.sub(r"^>\s*\*\*|\*\*", "", line) + "\n")
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_TEAL)
            p_pr.append(shd)
            index += 1
            continue

        if line.startswith("> "):
            paragraph = document.add_paragraph(line[2:].replace("**", ""))
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.right_indent = Cm(0.8)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(8)
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered and not in_references:
            paragraph = document.add_paragraph(style="Numbered Body")
            paragraph.add_run(f"{numbered.group(1)}. {numbered.group(2)}")
            index += 1
            continue

        if re.match(r"^\[\d+\]", line):
            document.add_paragraph(line, style="Reference Entry")
            index += 1
            continue

        if line.startswith("**Table"):
            document.add_paragraph(line.replace("**", ""), style="Table Caption")
            index += 1
            continue

        parts = [line]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if not nxt or nxt.startswith(("#", "|", ">", "![", "[[FIGURE_GRID|")) or re.match(r"^(\d+)\.\s+", nxt) or re.match(r"^\[\d+\]", nxt):
                break
            parts.append(nxt)
            index += 1
        add_body_paragraph(document, " ".join(parts))


def enable_field_updates(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> None:
    # Use the supplied graduation template as the package base so its official
    # page geometry, page borders, title-page behavior, header, footer, and theme
    # remain intact. The template lacks numbering and grid-table definitions, so
    # those are added programmatically after its empty body is cleared.
    document = Document(TEMPLATE)
    clear_body(document)
    style_document(document)
    num_id = configure_multilevel_numbering(document)
    add_cover(document)
    add_front_matter(document)
    for source in SOURCES:
        add_chapter_from_markdown(document, num_id, source)
    enable_field_updates(document)
    document.core_properties.title = "AscultiCor Graduation Project Book"
    document.core_properties.subject = "AI-powered ECG and PCG cardiac monitoring platform"
    document.core_properties.author = "AscultiCor Graduation Project Team"
    document.core_properties.keywords = "AscultiCor, ECG, PCG, IoT, AI, cardiac monitoring"
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
